"""Canva が書き出すような PPTX(全面背景 + テキスト + 画像 + 図形)をテスト用に作る。"""
import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

out = sys.argv[1]
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]
for n in range(2):
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(0xF3, 0xE8, 0xFF); bg.line.fill.background()
    tb = s.shapes.add_textbox(Emu(900000), Emu(600000), Emu(9000000), Emu(900000))
    tb.text_frame.text = f"タイトル {n+1}"; tb.text_frame.paragraphs[0].runs[0].font.size = Pt(44)
    tb2 = s.shapes.add_textbox(Emu(900000), Emu(2000000), Emu(6000000), Emu(900000))
    tb2.text_frame.text = "本文テキスト"; tb2.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(7500000), Emu(2000000), Emu(3000000), Emu(2000000))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x6D, 0x3D, 0xF5)
    pic = s.shapes.add_picture(sys.argv[2], Emu(900000), Emu(3500000), Emu(2500000), Emu(2500000))
prs.save(out)
